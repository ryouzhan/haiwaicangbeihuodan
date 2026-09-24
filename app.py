#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""海外仓备货单智能处理工具 (Web极简版) - 物流信息精准定位 + 在线商品库/密文同步双核驱动"""

from collections import defaultdict
from datetime import datetime
import io
import json
import os
import re
import sys
from typing import Any, Dict, List, Optional, Tuple
from cryptography.fernet import Fernet
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import pandas as pd
import streamlit as st

# ==================== 1. 页面配置与现代极简高级样式 ====================
st.set_page_config(
    page_title="海外仓备货发货处理工具",
    page_icon="📦",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown(
    """
<style>
    .stApp {
        background-color: #F8FAFC !important;
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
        color: #1E293B !important;
    }
    #MainMenu, footer, header {visibility: hidden;}

    .header-box { padding: 1.5rem 0 0.8rem 0; margin-bottom: 0.6rem; }
    .header-badge {
        display: inline-block; padding: 3px 10px; font-size: 0.72rem; font-weight: 600;
        text-transform: uppercase; background: #E2E8F0; color: #475569;
        border-radius: 9999px; margin-bottom: 0.3rem;
    }
    .header-title { font-size: 1.85rem; font-weight: 700; color: #0F172A; margin: 0; }
    .header-subtitle { font-size: 0.88rem; color: #64748B; margin-top: 0.2rem; }

    /* 顶部 Popover 胶囊按钮 (完全复刻 app.py 质感) */
    div[data-testid="stPopover"] > button {
        border-radius: 20px !important;
        padding: 4px 14px !important;
        font-size: 0.82rem !important;
        background: #FFFFFF !important;
        border: 1px solid #CBD5E1 !important;
        color: #334155 !important;
        height: auto !important;
        display: inline-flex !important;
        align-items: center !important;
        box-shadow: 0 1px 2px rgba(0,0,0,0.04) !important;
        transition: all 0.2s ease;
    }
    div[data-testid="stPopover"] > button:hover {
        border-color: #2563EB !important;
        color: #1E40AF !important;
        background: #F8FAFC !important;
    }

    [data-testid="stFileUploader"] section {
        background-color: #FFFFFF !important;
        border: 1.5px dashed #CBD5E1 !important;
        border-radius: 12px !important;
        padding: 1.5rem 1rem !important;
        box-shadow: 0 1px 3px rgba(0, 0, 0, 0.02) !important;
    }
    [data-testid="stFileUploader"] section:hover {
        border-color: #2563EB !important;
        background-color: #F8FAFC !important;
    }

    .metric-container {
        display: grid;
        grid-template-columns: repeat(6, 1fr);
        gap: 0.8rem;
        margin: 1.2rem 0 1.5rem 0;
    }
    .metric-card {
        background: #FFFFFF;
        padding: 1rem 0.8rem;
        border-radius: 10px;
        border: 1px solid #E2E8F0;
        box-shadow: 0 1px 2px rgba(0,0,0,0.03);
    }
    .metric-title { font-size: 0.75rem; font-weight: 500; color: #64748B; margin-bottom: 0.3rem; }
    .metric-num { font-size: 1.45rem; font-weight: 700; color: #0F172A; line-height: 1.1; }
    .metric-unit { font-size: 0.75rem; font-weight: 500; color: #94A3B8; margin-left: 0.2rem; }
</style>
""",
    unsafe_allow_html=True,
)


# ==================== 2. 商品库在线解密与特殊映射 (安全异常捕获) ====================
MAPPING_FILE = "sku_mapping.json"


def load_sku_mapping() -> Dict[str, str]:
  if os.path.exists(MAPPING_FILE):
    try:
      with open(MAPPING_FILE, "r", encoding="utf-8") as f:
        return json.load(f)
    except Exception:
      pass
  return {}


def save_sku_mapping(mapping: Dict[str, str]) -> None:
  try:
    with open(MAPPING_FILE, "w", encoding="utf-8") as f:
      json.dump(mapping, f, ensure_ascii=False, indent=2)
  except Exception:
    pass


def get_secret_key() -> Optional[bytes]:
  """获取解密密钥：优先 Streamlit Secrets，其次本地 secret.key (安全捕获异常)"""
  try:
    if "COMMODITIES_KEY" in st.secrets:
      return st.secrets["COMMODITIES_KEY"].encode()
  except Exception:
    pass

  if os.path.exists("secret.key"):
    try:
      with open("secret.key", "rb") as f:
        return f.read().strip()
    except Exception:
      pass
  return None


def parse_raw_table_bytes(raw_bytes: bytes) -> Optional[pd.DataFrame]:
  """从二进制字节流自动识别 Excel 或 CSV 并转为 DataFrame"""
  try:
    if raw_bytes.startswith(b"PK\x03\x04") or raw_bytes.startswith(
        b"\xd0\xcf\x11\xe0"
    ):
      return pd.read_excel(io.BytesIO(raw_bytes))
    else:
      try:
        return pd.read_csv(io.BytesIO(raw_bytes), encoding="utf-8-sig")
      except Exception:
        return pd.read_csv(io.BytesIO(raw_bytes), encoding="gbk")
  except Exception:
    return None


def load_active_commodities() -> Tuple[Optional[pd.DataFrame], str]:
  """智能定位并加载商品表：

  1. 优先解密云端/本地 commodities.dat
  2. 若无则检索本地明文表格
  """
  key = get_secret_key()

  if os.path.exists("commodities.dat"):
    if not key:
      return None, "未配置解密密钥 (请在 Secrets 填入 COMMODITIES_KEY)"
    try:
      with open("commodities.dat", "rb") as f:
        cipher_data = f.read()
      cipher = Fernet(key)
      decrypted_bytes = cipher.decrypt(cipher_data)
      df = parse_raw_table_bytes(decrypted_bytes)
      if df is not None:
        df.columns = [str(c).strip() for c in df.columns]
        return df, "商品库 (已加密安全同步)"
    except Exception as e:
      return None, f"解密失败: {e}"

  # 检索本地明文表格
  valid_exts = (".xlsx", ".xls", ".csv")
  candidates = []
  for fname in os.listdir("."):
    if fname.startswith("~$"):
      continue
    if any(fname.lower().endswith(ext) for ext in valid_exts) and "commodit" in fname.lower():
      full_path = os.path.join(".", fname)
      candidates.append((full_path, fname, os.path.getmtime(full_path)))

  if candidates:
    candidates.sort(key=lambda x: x, reverse=True)
    latest_path, latest_name, _ = candidates[0]
    try:
      with open(latest_path, "rb") as f:
        df = parse_raw_table_bytes(f.read())
      if df is not None:
        df.columns = [str(c).strip() for c in df.columns]
        return df, latest_name
    except Exception:
      pass

  return None, "未检测到商品库"


# ==================== 3. 复合表格解析与单号精准提取 ====================
def parse_raw_order_file(uploaded_file):
  """精准拆解复合表格：

  1. 截取有效商品明细
  2. 定位【物流信息】区块中的【关联备货单号】作为货件编号
  """
  filename = str(uploaded_file.name).lower()
  if filename.endswith(".csv"):
    df_raw = pd.read_csv(uploaded_file, header=None)
  else:
    df_raw = pd.read_excel(uploaded_file, header=None)

  # 寻找有效商品行截断点
  cut_idx = len(df_raw)
  split_keywords = ["备货单号", "辅料SKU", "关联备货单号", "三方仓入库单号"]

  for idx in range(1, len(df_raw)):
    first_cell = str(df_raw.iloc[idx, 0]).strip()
    if any(k in first_cell for k in split_keywords):
      cut_idx = idx
      break

  goods_headers = [str(c).strip() for c in df_raw.iloc[0].values]
  goods_df = df_raw.iloc[1:cut_idx].copy()
  goods_df.columns = goods_headers
  goods_df = goods_df.loc[:, ~goods_df.columns.str.startswith("Unnamed")]
  goods_df.dropna(how="all", inplace=True)

  # 从【物流信息】中定位【关联备货单号】
  related_order_code = ""

  for row_idx in range(cut_idx, len(df_raw)):
    row_vals = [str(v).strip() for v in df_raw.iloc[row_idx].values]
    if "关联备货单号" in row_vals:
      col_idx = row_vals.index("关联备货单号")
      if row_idx + 1 < len(df_raw):
        val = str(df_raw.iloc[row_idx + 1, col_idx]).strip()
        if val and val != "nan" and val != "None":
          related_order_code = val
          break

  # 兜底正则提取 OWS / FBA
  if not related_order_code:
    for row_idx in range(cut_idx, len(df_raw)):
      for cell in df_raw.iloc[row_idx].dropna():
        s = str(cell).strip()
        m = re.search(r"\b(OWS[A-Za-z0-9\-]+|FBA[A-Za-z0-9\-]+)\b", s)
        if m:
          related_order_code = m.group(1)
          break
      if related_order_code:
        break

  return goods_df, (related_order_code or "未知单号")


def get_region(addr):
  addr = str(addr).upper()
  if "AWD" in addr:
    return "AWD仓"
  east = ["ABE8", "AVP1", "DCA6", "TEB9", "PHL7", "BDL3"]
  central = ["DFW6", "FOE1", "MDW2", "IND7", "MEM1"]
  west = ["IUTE", "ONT8", "LAS1", "LAX9", "GYR2", "ABQ2", "SCK4", "SMF3"]
  if any(x in addr for x in east):
    return "美东"
  if any(x in addr for x in central):
    return "美中"
  if any(x in addr for x in west):
    return "美西"
  return "其他"


def extract_pcs_from_text(text):
  match = re.search(r"(\d+)\s*(?:pc|pcs|PC|PCS|只|件|套)", str(text), re.I)
  if match:
    return int(match.group(1))
  return 1


# ==================== 4. 核心计算与商品属性合并 ====================
def process_shipment_data(
    goods_df, order_code, commodities_df, sku_mapping=None
):
  """基于商品库匹配长宽高、单箱重量、单价与PCS，完成全维度核算"""
  goods_df["SKU"] = goods_df["SKU"].astype(str).str.strip()

  comm_dict = {}
  if commodities_df is not None and not commodities_df.empty:
    sku_col = next(
        (c for c in commodities_df.columns if str(c).strip().upper() == "SKU"),
        None,
    )
    if sku_col:
      for _, r in commodities_df.iterrows():
        s = str(r[sku_col]).strip()
        comm_dict[s] = r.to_dict()

  merged_rows = []
  missing_skus = []

  for _, row in goods_df.iterrows():
    raw_sku = str(row["SKU"]).strip()

    # 应用特殊 SKU 别名映射
    lookup_sku = (
        sku_mapping.get(raw_sku, raw_sku) if sku_mapping else raw_sku
    )

    matched = comm_dict.get(lookup_sku)
    if not matched and lookup_sku.upper().startswith("ZF-"):
      matched = comm_dict.get(lookup_sku[3:])
    if not matched and not lookup_sku.upper().startswith("ZF-"):
      matched = comm_dict.get(f"ZF-{lookup_sku}")
    if not matched:
      matched = comm_dict.get(raw_sku, {})

    if not matched:
      missing_skus.append(raw_sku)

    item = {**row.to_dict()}

    # 1. 品名与品牌
    item["_品名"] = item.get("品名") or matched.get(
        "品名", matched.get("中文品名", "")
    )
    item["_供应商"] = matched.get(
        "供应商名称", matched.get("供应商", matched.get("商品品牌", ""))
    )

    # 2. 单箱数量（箱规：一箱几套）
    local_carton = row.get("单箱数量") or row.get("单箱数量(pcs)") or 0
    cloud_carton = matched.get("单箱数量(pcs)", matched.get("单箱数量", 0))
    final_carton = pd.to_numeric(
        local_carton if local_carton and local_carton != 0 else cloud_carton,
        errors="coerce",
    )
    item["_单箱数量"] = (
        int(final_carton)
        if not np.isnan(final_carton) and final_carton > 0
        else 0
    )

    # 3. 单品 PCS（每套包含几个单件）
    pcs_val = pd.to_numeric(
        matched.get("单品PCS", matched.get("PCS", 0)), errors="coerce"
    )
    if not np.isnan(pcs_val) and pcs_val > 0:
      item["_单品PCS"] = int(pcs_val)
    else:
      item["_单品PCS"] = extract_pcs_from_text(item["_品名"])

    # 4. 采购单价
    price_val = 0.0
    for p_col in [
        "采购单价(CNY)",
        "采购成本(￥)",
        "指定采购单价(CNY)",
        "采购单价",
        "单价",
    ]:
      if p_col in matched:
        val = pd.to_numeric(matched[p_col], errors="coerce")
        if not np.isnan(val) and val > 0:
          price_val = val
          break
    item["_单价"] = price_val

    # 5. 箱规长宽高与重量
    def get_dim(keys, fallback=0.0):
      for k in keys:
        if k in matched:
          v = pd.to_numeric(matched[k], errors="coerce")
          if not np.isnan(v) and v > 0:
            return float(v)
      return float(fallback)

    item["_重量"] = get_dim(
        ["单箱重量(kg)", "重量(kg)", "毛重(kg)", "单箱重量"],
        row.get("单箱重量(kg)", 0),
    )
    item["_长"] = get_dim(
        ["箱规长(cm)", "外箱长(cm)", "长(cm)", "长"], row.get("箱规长(cm)", 0)
    )
    item["_宽"] = get_dim(
        ["箱规宽(cm)", "外箱宽(cm)", "宽(cm)", "宽"], row.get("箱规宽(cm)", 0)
    )
    item["_高"] = get_dim(
        ["箱规高(cm)", "外箱高(cm)", "高(cm)", "高"], row.get("箱规高(cm)", 0)
    )

    merged_rows.append(item)

  df_m = pd.DataFrame(merged_rows)

  df_detail = pd.DataFrame()
  df_detail["SKU"] = df_m["SKU"]
  df_detail["品名"] = df_m["_品名"]
  df_detail["货件编号"] = order_code

  df_detail["单箱数量"] = df_m["_单箱数量"]
  df_detail["单品PCS"] = df_m["_单品PCS"]

  if "箱数" in df_m.columns and "备货量" in df_m.columns:
    df_detail["箱数"] = (
        pd.to_numeric(df_m["箱数"], errors="coerce").fillna(0).astype(int)
    )
    df_detail["备货套数"] = (
        pd.to_numeric(df_m["备货量"], errors="coerce").fillna(0).astype(int)
    )
    df_detail["备货套数"] = np.where(
        df_detail["备货套数"] > 0,
        df_detail["备货套数"],
        df_detail["箱数"] * df_detail["单箱数量"],
    )
  else:
    raw_qty = (
        pd.to_numeric(
            df_m.get("发货量", df_m.get("申报量", df_m.get("备货量", 0))),
            errors="coerce",
        )
        .fillna(0)
        .astype(int)
    )
    safe_carton = np.where(
        df_detail["单箱数量"] > 0, df_detail["单箱数量"], 1
    )
    df_detail["箱数"] = np.where(
        df_detail["单箱数量"] > 0,
        np.ceil(raw_qty / safe_carton).astype(int),
        0,
    )
    df_detail["备货套数"] = raw_qty

  df_detail["总PCS"] = df_detail["备货套数"] * df_detail["单品PCS"]

  df_detail["单箱重量(kg)"] = df_m["_重量"].round(2)
  df_detail["外箱总重量(kg)"] = (
      df_detail["单箱重量(kg)"] * df_detail["箱数"]
  ).round(2)
  df_detail["箱规长(cm)"] = df_m["_长"].round(2)
  df_detail["箱规宽(cm)"] = df_m["_宽"].round(2)
  df_detail["箱规高(cm)"] = df_m["_高"].round(2)

  vol = (df_m["_长"] * df_m["_宽"] * df_m["_高"]) / 1000000
  df_detail["外箱总体积(m³)"] = (vol * df_detail["箱数"]).round(3)
  df_detail["外箱总体积重(kg)"] = (df_detail["外箱总体积(m³)"] * 167).round(2)

  df_detail["采购单价"] = df_m["_单价"].round(2)
  df_detail["总货值(￥)"] = (
      df_detail["备货套数"] * df_detail["采购单价"]
  ).round(2)
  df_detail["供应商"] = df_m["_供应商"]

  addr = ""
  for col in ["配送地址", "物流中心编码", "收货仓库"]:
    if col in df_m.columns:
      addr = df_m[col].iloc[0]
      break
  df_detail["配送地址"] = addr or "AWD仓"
  df_detail["仓库分区"] = df_detail["配送地址"].apply(get_region)

  summary = (
      df_detail.groupby("货件编号")
      .agg({
          "仓库分区": "first",
          "配送地址": "first",
          "箱数": "sum",
          "备货套数": "sum",
          "总PCS": "sum",
          "外箱总重量(kg)": "sum",
          "外箱总体积(m³)": "sum",
          "外箱总体积重(kg)": "sum",
          "总货值(￥)": "sum",
      })
      .reset_index()
  )

  summary.rename(
      columns={
          "箱数": "总箱数",
          "备货套数": "总套数",
          "总货值(￥)": "货件总货值(￥)",
      },
      inplace=True,
  )

  return df_detail, summary, list(set(missing_skus))


# ==================== 5. 专业 Excel 导出与自动美化 ====================
def export_and_beautify(df_detail, df_summary):
  output = io.BytesIO()
  with pd.ExcelWriter(output, engine="openpyxl") as writer:
    df_summary.to_excel(writer, sheet_name="汇总结果", index=False)
    df_detail.to_excel(writer, sheet_name="详细数据", index=False)

  wb = load_workbook(output)
  header_fill = PatternFill(
      start_color="1E293B", end_color="1E293B", fill_type="solid"
  )
  header_font = Font(color="FFFFFF", bold=True, name="微软雅黑", size=10)
  stripe_fill = PatternFill(
      start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
  )
  thin_border = Border(
      left=Side(style="thin", color="CBD5E1"),
      right=Side(style="thin", color="CBD5E1"),
      top=Side(style="thin", color="CBD5E1"),
      bottom=Side(style="thin", color="CBD5E1"),
  )
  center_align = Alignment(
      horizontal="center", vertical="center", wrap_text=True
  )

  for sheetname in ["汇总结果", "详细数据"]:
    ws = wb[sheetname]
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row)):
      for cell in row:
        cell.border = thin_border
        cell.alignment = center_align
        if i == 0:
          cell.fill = header_fill
          cell.font = header_font
        elif i % 2 == 0:
          cell.fill = stripe_fill

    for col in ws.columns:
      max_len = 0
      col_letter = col[0].column_letter
      for cell in col:
        try:
          length = len(str(cell.value).encode("gbk"))
          if length > max_len:
            max_len = length
        except:
          pass
      ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

  final_stream = io.BytesIO()
  wb.save(final_stream)
  return final_stream.getvalue()


# ==================== 6. 主程序与界面交互 ====================
def main():
  st.markdown(
      """
    <div class="header-box">
        <div class="header-badge">AWD Cloud Engine V9.1</div>
        <h1 class="header-title">海外仓备货发货单智能处理</h1>
        <p class="header-subtitle">物流信息单号联动 · 在线安全商品库同步 · 箱规重量体积货值全核算</p>
    </div>
    """,
      unsafe_allow_html=True,
  )

  # 1. 特殊映射字典
  if "sku_mapping" not in st.session_state:
    st.session_state["sku_mapping"] = load_sku_mapping()
  current_mapping = st.session_state["sku_mapping"]

  # 2. 自动定位商品库
  active_df, table_label = load_active_commodities()

  custom_uploaded = st.session_state.get("custom_commodities", None)
  if custom_uploaded is not None:
    table_pill_label = f"🟢 自定义: {custom_uploaded.name} ▾"
  elif active_df is not None:
    table_pill_label = f"🟢 {table_label} ▾"
  else:
    table_pill_label = f"🔴 {table_label} ▾"

  map_count = len(current_mapping)
  mapping_pill_label = (
      f"⚡ 特殊映射 ({map_count}条) ▾" if map_count > 0 else "⚡ 特殊映射 ▾"
  )

  # 3. 顶部胶囊组件
  col_p1, col_p2, _ = st.columns([1.5, 1.2, 1.3])

  with col_p1:
    with st.popover(table_pill_label):
      st.caption("临时更换商品库（仅本次生效）：")
      custom_file = st.file_uploader(
          "上传替代商品列表",
          type=["xlsx", "xls", "csv"],
          label_visibility="collapsed",
          key="custom_commodities",
      )
      if custom_uploaded is not None and st.button(
          "恢复默认商品库", use_container_width=True
      ):
        del st.session_state["custom_commodities"]
        st.rerun()

  with col_p2:
    with st.popover(mapping_pill_label):
      st.caption("双击编辑，支持从 Excel 复制两列直接粘贴：")
      rows = [
          {"面单SKU": k, "商品库SKU": v} for k, v in current_mapping.items()
      ]
      if not rows:
        rows = [{"面单SKU": "", "商品库SKU": ""}]
      df_mapping = pd.DataFrame(rows)

      edited_df = st.data_editor(
          df_mapping,
          num_rows="dynamic",
          use_container_width=True,
          hide_index=True,
          height=200,
          column_config={
              "面单SKU": st.column_config.TextColumn(
                  "备货单 SKU", required=True
              ),
              "商品库SKU": st.column_config.TextColumn(
                  "商品库 SKU", required=True
              ),
          },
          key="sku_mapping_editor",
      )

      c_btn1, c_btn2 = st.columns(2)
      if c_btn1.button("保存规则", type="primary", use_container_width=True):
        new_map = {}
        for _, r in edited_df.iterrows():
          src = str(r.get("面单SKU", "")).strip()
          tgt = str(r.get("商品库SKU", "")).strip()
          if src and tgt and src != "nan" and tgt != "nan":
            new_map[src] = tgt
        st.session_state["sku_mapping"] = new_map
        save_sku_mapping(new_map)
        st.rerun()

      if c_btn2.button("清空全部", use_container_width=True):
        st.session_state["sku_mapping"] = {}
        save_sku_mapping({})
        st.rerun()

  if custom_uploaded is not None:
    active_df = parse_raw_table_bytes(custom_uploaded.getvalue())

  st.write("")

  # 4. 主发货单上传
  uploaded_file = st.file_uploader(
      "请上传备货单/发货单 Excel 或 CSV 文件",
      type=["xlsx", "xls", "csv"],
      help="自动识别商品行并从【物流信息】中定位【关联备货单号】",
  )

  if uploaded_file is not None:
    try:
      with st.spinner("正在解析物流信息并匹配商品库数据..."):
        goods_df, related_order_code = parse_raw_order_file(uploaded_file)
        df_detail, df_summary, missing_skus = process_shipment_data(
            goods_df, related_order_code, active_df, sku_mapping=current_mapping
        )

      total_box = int(df_summary["总箱数"].sum())
      total_sets = int(df_summary["总套数"].sum())
      total_pcs = int(df_summary["总PCS"].sum())
      total_wt = f"{df_summary['外箱总重量(kg)'].sum():,.2f}"
      total_vol = f"{df_summary['外箱总体积(m³)'].sum():,.2f}"
      total_val = f"￥{df_summary['货件总货值(￥)'].sum():,.2f}"

      st.markdown(
          f"""
            <div class="metric-container">
                <div class="metric-card">
                    <div class="metric-title">货件编号 (关联备货单)</div>
                    <div class="metric-num" style="font-size:1.05rem; word-break:break-all;">{related_order_code}</div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">总装箱量</div>
                    <div class="metric-num">{total_box:,}<span class="metric-unit">箱</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">总备货量</div>
                    <div class="metric-num">{total_sets:,}<span class="metric-unit">套</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">总 PCS (单件实物)</div>
                    <div class="metric-num">{total_pcs:,}<span class="metric-unit">件</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">实重 / 总体积</div>
                    <div class="metric-num" style="font-size:1.15rem;">{total_wt}<span class="metric-unit">kg</span> / {total_vol}<span class="metric-unit">m³</span></div>
                </div>
                <div class="metric-card">
                    <div class="metric-title">货件总货值</div>
                    <div class="metric-num">{total_val}</div>
                </div>
            </div>
            """,
          unsafe_allow_html=True,
      )

      if missing_skus:
        st.warning(
            f"⚠️ 提示：共有 {len(missing_skus)} 个 SKU 未在商品库中匹配到规格信息："
            f" {', '.join(missing_skus[:8])}{'...' if len(missing_skus) > 8 else ''}"
        )

      ts = datetime.now().strftime("%Y%m%d_%H%M%S")
      excel_bytes = export_and_beautify(df_detail, df_summary)

      st.download_button(
          label="⬇️ 导出全量发货单与汇总表 (.xlsx)",
          data=excel_bytes,
          file_name=f"发货单处理结果_{related_order_code}_{ts}.xlsx",
          mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
          type="primary",
          use_container_width=True,
      )

      tab1, tab2 = st.tabs(["📊 货件汇总表", "📝 计算明细表"])
      with tab1:
        st.dataframe(df_summary, use_container_width=True)
      with tab2:
        st.dataframe(df_detail, use_container_width=True)

    except Exception as e:
      st.error(f"❌ 数据处理失败：{str(e)}")


if __name__ == "__main__":
  main()